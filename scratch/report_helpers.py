import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Đặt màu nền cho một ô trong bảng."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Đặt lề trong cho ô của bảng (đơn vị dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, hex_color="CCCCCC", sz="4", val="single"):
    """Thiết lập viền mỏng tinh tế cho toàn bảng."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{hex_color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, text_list, title="LƯU Ý QUAN TRỌNG:"):
    """Tạo hộp ghi chú chuyên nghiệp có viền trái nổi bật."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Border: thick left border, no top/bottom/right
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"📌 {title}\n")
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    for item in text_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15
        run_i = p_item.add_run(f"• {item}")
        run_i.font.name = "Times New Roman"
        run_i.font.size = Pt(10.5)
        run_i.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    doc.add_paragraph() # Spacing

def add_code_block(doc, code_str, caption=""):
    """Tạo hộp code snippet với phông chữ monospace."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>\n'
        f'  <w:left w:val="single" w:sz="18" w:space="0" w:color="4A5568"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
    
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(6)
        run_cap = p_cap.add_run(f"Đoạn mã: {caption}")
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    else:
        doc.add_paragraph()

print("Helper definitions loaded successfully.")
